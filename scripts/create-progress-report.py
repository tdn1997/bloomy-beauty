"""Build the Vietnamese first progress report from freshly captured screenshots."""
from pathlib import Path
from datetime import datetime
import json, hashlib
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs/bao-cao/tien-do-lan-1'
MEMBERS=[('Phạm Văn Bình','25TX610001'),('Nguyễn Duy Thuận','25TX810025'),('Đoàn Thị Mỹ Thiện','25TX810023')]
SCREENS=[
('01-dang-nhap','Đăng nhập','Người dùng nhập email và mật khẩu để truy cập ứng dụng.'),
('02-loi-dang-nhap','Kiểm tra dữ liệu đăng nhập','Gửi biểu mẫu trống: ứng dụng hiển thị thông báo yêu cầu dữ liệu hợp lệ.'),
('03-dang-ky','Đăng ký tài khoản','Biểu mẫu gồm họ tên, email, mật khẩu và xác nhận mật khẩu.'),
('04-trang-chu','Trang chủ sau đăng ký','Đăng ký tài khoản kiểm thử thành công; hiển thị lời chào, bộ sưu tập và sản phẩm.'),
('05-gio-hang-trong','Giỏ hàng trống','Tài khoản mới chưa có sản phẩm; tổng tiền bằng 0 và chưa thể tiếp tục đặt hàng.'),
('06-danh-muc','Khám phá sản phẩm','Duyệt danh mục, bộ lọc và lựa chọn sắp xếp sản phẩm.'),
('07-tim-kiem-rong','Không có kết quả tìm kiếm','Từ khóa không khớp: hiển thị 0 sản phẩm và cho phép xóa bộ lọc.'),
('08-yeu-thich','Danh sách yêu thích','Đánh dấu gel bí đao rửa mặt và xem lại ở thẻ Yêu thích.'),
('09-chi-tiet-san-pham','Chi tiết sản phẩm','Xem ảnh, thương hiệu, dung tích và giá của mặt hàng đã chọn.'),
('10-chi-tiet-noi-dung','Thông tin và thao tác sản phẩm','Xem tồn kho, mô tả, thành phần, cách dùng và nút Thêm vào giỏ.'),
('11-gio-hang','Giỏ hàng có sản phẩm','Một gel bí đao rửa mặt giá 192.000 ₫; có điều chỉnh số lượng và xóa.'),
('12-giao-hang','Nhập thông tin giao hàng','Cung cấp người nhận, số điện thoại và địa chỉ trước khi đặt COD.'),
('13-loi-giao-hang','Kiểm tra thông tin giao hàng','Bỏ trống thông tin liên hệ: ứng dụng giữ biểu mẫu và báo lỗi để bổ sung.'),
('14-xac-nhan-cod','Kiểm tra trước khi đặt COD','Đối chiếu thông tin người nhận, hàng hóa, phí giao hàng và tổng thanh toán.'),
('15-dat-hang-thanh-cong','Đặt hàng thành công','Hệ thống tạo mã đơn, lưu đơn và hiển thị kết quả đặt COD.'),
('16-danh-sach-don','Đơn hàng của tôi','Tài khoản kiểm thử xem đơn vừa tạo với trạng thái Chờ xác nhận.'),
('17-chi-tiet-don','Chi tiết đơn chờ xác nhận','Xem hàng hóa, địa chỉ, tổng tiền và quyền hủy khi đơn còn chờ.'),
('18-tai-khoan','Thông tin tài khoản','Mở hồ sơ, lịch sử đơn hàng và chức năng đăng xuất.'),
('19-chinh-sua-ho-so','Chỉnh sửa hồ sơ','Biểu mẫu họ tên và địa chỉ mặc định; ảnh ghi nhận giao diện, chưa xác nhận lưu thay đổi.'),
('20-quan-tri-san-pham','Quản trị sản phẩm','Quản trị viên xem danh sách hàng hóa, giá và tồn kho hiện tại.'),
('21-them-san-pham','Biểu mẫu thêm sản phẩm','Giao diện khai báo mặt hàng mới; lượt chụp này không lưu sản phẩm mới.'),
('22-quan-tri-don','Quản trị đơn hàng','Quản trị viên mở danh sách đơn và lựa chọn đơn kiểm thử vừa tạo.'),
('23-admin-cho-xac-nhan','Đơn chờ xử lý','Chi tiết đơn và các thao tác chuyển trạng thái dành cho quản trị viên.'),
('24-xac-nhan-cap-nhat','Xác nhận cập nhật trạng thái','Hộp thoại yêu cầu xác nhận trước khi lưu chuyển trạng thái đơn.'),
('25-da-xac-nhan','Đã xác nhận','Đơn chuyển từ Chờ xác nhận sang Đã xác nhận; lưu lịch sử xử lý.'),
('26-dang-giao','Đang giao','Quản trị viên chuyển đơn sang Đang giao theo luồng nghiệp vụ.'),
('27-da-giao-admin','Đã giao – phía quản trị','Hoàn tất chuỗi xử lý đơn và hiển thị lịch sử chuyển trạng thái.'),
('28-da-giao-khach','Đã giao – phía khách hàng','Khách đăng nhập lại và thấy trạng thái Đã giao; không còn thao tác hủy.'),
('29-xac-nhan-huy','Xác nhận hủy đơn khác','Nhánh kiểm thử bổ sung: tạo một đơn khác rồi yêu cầu hủy khi còn chờ xác nhận.'),
('30-da-huy','Đơn đã hủy','Đơn thứ hai chuyển sang Đã hủy; đơn đã giao trước đó vẫn giữ nguyên trạng thái.'),
]

def table(doc,headers,rows,widths=None):
    t=doc.add_table(rows=1, cols=len(headers));t.style='Table Grid'
    for c,s in zip(t.rows[0].cells,headers):
        c.text=s
        for r in c.paragraphs[0].runs:r.bold=True
        pr=c._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F3E5EB');pr.append(shade)
    for row in rows:
        for c,s in zip(t.add_row().cells,row):c.text=s
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(5)
                for r in p.runs:r.font.size=Pt(10)
    return t

def make():
    missing=[key for key,_,_ in SCREENS if not (OUT/(key+'.png')).exists()]
    if missing:raise SystemExit('Missing screenshots: '+', '.join(missing))
    d=Document();s=d.sections[0];s.page_height=Cm(29.7);s.page_width=Cm(21)
    s.top_margin=Cm(1.8);s.bottom_margin=Cm(1.7);s.left_margin=Cm(2);s.right_margin=Cm(2)
    normal=d.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12)
    normal.paragraph_format.line_spacing=1.15;normal.paragraph_format.space_after=Pt(7)
    for style in ['Title','Heading 1','Heading 2','Heading 3']:
        d.styles[style].font.name='Times New Roman';d.styles[style].font.color.rgb=RGBColor.from_string('803C59')
    h=s.header.paragraphs[0];h.text='NHÓM 4  |  BLOOMY BEAUTY  |  BÁO CÁO TIẾN ĐỘ LẦN 1';h.runs[0].font.size=Pt(9)
    foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    foot.add_run('Nhóm 4 • Tiến độ lần 1 • Trang ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
    for r in foot.runs:r.font.size=Pt(9)
    d.add_paragraph('MÔN HỌC: LẬP TRÌNH DI ĐỘNG').alignment=1
    d.add_paragraph('\n\n')
    d.add_heading('BÁO CÁO TIẾN ĐỘ\nLẦN 1',0).alignment=1
    d.add_paragraph('ĐỀ TÀI: ỨNG DỤNG BÁN MỸ PHẨM\nBLOOMY BEAUTY',style='Subtitle').alignment=1
    d.add_paragraph('NHÓM 4').alignment=1
    table(d,['Họ và tên','Mã số sinh viên'],MEMBERS)
    d.add_paragraph('\nNgày báo cáo: 14/10/2026\nTrạng thái chung: Đang kiểm thử và hoàn thiện').alignment=1
    d.add_paragraph('Ảnh giao diện được chụp trực tiếp trên máy giả lập Android API 36 trong lượt thực hiện báo cáo.').alignment=1
    d.add_page_break()
    d.add_heading('1. Mục tiêu và phạm vi báo cáo',1)
    d.add_paragraph('Bloomy Beauty là ứng dụng Android hỗ trợ xem và tìm mỹ phẩm, lưu yêu thích, quản lý giỏ hàng, đặt hàng COD, theo dõi đơn và quản trị danh mục. Báo cáo lần 1 ghi nhận hiện trạng triển khai và các màn hình theo trình tự sử dụng tại ngày báo cáo.')
    d.add_paragraph('Công nghệ hiện có: Kotlin, Jetpack Compose, Material 3, ViewModel, Coroutines/StateFlow và SQLite schema 5. Dữ liệu tài khoản, sản phẩm, giỏ hàng và đơn hàng được lưu cục bộ trên một thiết bị. Bản demo có 8 sản phẩm tham khảo Cocoon với ảnh đóng gói ngoại tuyến.')
    d.add_heading('Quy ước trạng thái',2)
    table(d,['Trạng thái','Ý nghĩa'],[
        ('Đã triển khai','Chức năng và giao diện đã có trong bản demo hiện tại.'),
        ('Đã kiểm tra trong lượt chụp','Đã thực hiện thao tác tương ứng và ghi nhận ảnh trong báo cáo này.'),
        ('Đang kiểm thử','Tiếp tục rà soát đầu vào, trường hợp biên, lưu trữ và thiết bị; chưa coi là nghiệm thu toàn bộ.'),
        ('Chưa triển khai','Chức năng mở rộng chưa có trong phạm vi hiện tại.')])
    d.add_heading('Trình tự sử dụng được ghi nhận',2)
    d.add_paragraph('Đăng nhập/đăng ký → trang chủ → danh mục/tìm kiếm → yêu thích → chi tiết → giỏ hàng → giao hàng → xác nhận COD → đặt thành công → lịch sử/chi tiết đơn → hồ sơ → đăng nhập quản trị → quản lý sản phẩm → xử lý đơn → khách kiểm tra Đã giao. Nhánh bổ sung: tạo đơn khác và hủy khi còn Chờ xác nhận.')
    d.add_paragraph('Phần hồ sơ và thêm sản phẩm được ghi nhận ở mức giao diện trong lượt chụp. Không quy kết các ảnh biểu mẫu này là bằng chứng lưu dữ liệu thành công.')
    d.add_page_break()
    d.add_heading('2. Tiến độ thực hiện lần 1',1)
    table(d,['Hạng mục','Hiện trạng','Theo dõi tiếp'],[
      ('Tài khoản','Đã triển khai; đã đăng ký tài khoản kiểm thử, đăng nhập và đăng xuất.','Đang kiểm thử dữ liệu không hợp lệ và phiên đăng nhập.'),
      ('Danh mục và tìm kiếm','Đã triển khai; đã xem danh mục, sản phẩm và kết quả rỗng.','Đang kiểm thử lọc, sắp xếp và nhập tiếng Việt.'),
      ('Chi tiết và yêu thích','Đã triển khai; đã thêm yêu thích và xem nội dung sản phẩm.','Tiếp tục rà soát nội dung và bố cục.'),
      ('Giỏ hàng và giao hàng','Đã triển khai; đã kiểm tra giỏ rỗng, giỏ có hàng và dữ liệu giao hàng thiếu.','Đang kiểm thử số lượng, giá và tồn kho thay đổi.'),
      ('Đặt COD','Đã triển khai; đã tạo đơn và xem kết quả thực tế.','Đang kiểm thử gửi lặp, khôi phục và nhất quán dữ liệu.'),
      ('Theo dõi và hủy đơn','Đã triển khai; đã xem đơn Đã giao và kiểm tra hủy đơn chờ.','Tiếp tục kiểm thử quyền sở hữu và trạng thái không hợp lệ.'),
      ('Hồ sơ','Đã triển khai; đã chụp biểu mẫu chỉnh sửa.','Đang kiểm thử lưu và đọc lại thông tin.'),
      ('Quản trị','Đã triển khai; đã xem danh mục, biểu mẫu thêm và chuyển trạng thái đơn.','Đang kiểm thử thêm/sửa/ẩn sản phẩm và phân quyền.'),
      ('Tài liệu và minh chứng','Đã bổ sung ảnh theo luồng và báo cáo Word lần 1.','Hoàn thiện thông tin học phần theo yêu cầu nộp.'),
      ('Backend, thanh toán online','Chưa triển khai; COD và kho là dữ liệu demo cục bộ.','Xem xét ở giai đoạn mở rộng.')])
    d.add_paragraph('Đây là báo cáo tiến độ, chưa phải biên bản nghiệm thu cuối cùng. Chưa có xác nhận phân công và tỷ lệ đóng góp thực tế, vì vậy báo cáo chỉ ghi danh sách thành viên, không tự gán công việc đã làm cho từng người.')
    d.add_page_break()
    d.add_heading('3. Minh chứng màn hình theo thứ tự sử dụng',1)
    d.add_paragraph('Danh mục tra cứu: số hình trùng số bước. Tất cả ảnh trong mục này là ảnh chụp mới từ bản demo; dữ liệu người nhận và đơn hàng dùng cho kiểm thử.')
    for i,(key,title,desc) in enumerate(SCREENS,1):
        d.add_paragraph(f'{i:02d}. {title}',style=None).paragraph_format.space_after=Pt(2)
    for offset in range(0,len(SCREENS),2):
        d.add_page_break()
        d.add_heading(f'3.{offset//2+1}. Các bước {offset+1:02d}–{min(offset+2,len(SCREENS)):02d}',1)
        t=d.add_table(rows=1,cols=2);t.autofit=False
        for c in t.columns:c.width=Cm(8.3)
        for j,(key,title,desc) in enumerate(SCREENS[offset:offset+2]):
            cell=t.rows[0].cells[j]
            p=cell.paragraphs[0];r=p.add_run(f'Hình {offset+j+1:02d}. {title}');r.bold=True;r.font.size=Pt(11)
            rowpr=t.rows[0]._tr.get_or_add_trPr()
            if rowpr.find(qn('w:cantSplit')) is None:rowpr.append(OxmlElement('w:cantSplit'))
            p=cell.add_paragraph();p.alignment=1;p.add_run().add_picture(str(OUT/(key+'.png')),height=Cm(14.3))
            p=cell.add_paragraph(desc)
            for r in p.runs:r.font.size=Pt(10)
            p=cell.add_paragraph('Trạng thái: '+('Đã ghi nhận giao diện; đang kiểm thử lưu dữ liệu.' if key.startswith(('19-','21-')) else 'Đã ghi nhận trong lượt kiểm tra; tiếp tục kiểm thử.'))
            for r in p.runs:r.font.size=Pt(10);r.italic=True
    d.add_page_break()
    d.add_heading('4. Kết quả kiểm tra và công việc tiếp theo',1)
    d.add_paragraph('Lượt kiểm tra ngày 14/10/2026 tập trung vào thao tác giao diện trên bản demo: tài khoản mới, trạng thái rỗng/lỗi đầu vào, đặt COD, chuyển trạng thái đơn và hủy đơn chờ. Các trạng thái đơn được minh họa bằng thao tác thật trên dữ liệu kiểm thử.')
    table(d,['Nội dung','Kết quả ghi nhận'],[
      ('Luồng đặt hàng','Một gel bí đao: 192.000 ₫ + phí giao hàng 30.000 ₫ = 222.000 ₫ COD.'),
      ('Luồng xử lý đơn','Chờ xác nhận → Đã xác nhận → Đang giao → Đã giao; khách xem được trạng thái cuối.'),
      ('Nhánh hủy','Tạo đơn khác ở Chờ xác nhận, xác nhận hủy và xem trạng thái Đã hủy.'),
      ('Trạng thái kiểm thử','Có ảnh đăng nhập thiếu dữ liệu, giỏ hàng trống, tìm kiếm rỗng và giao hàng thiếu thông tin.'),
      ('Kết quả tự động lưu sẵn','Hồ sơ trong deliverables/ket-qua-kiem-thu: 24 JVM và 81 Android, không có ca thất bại. Đây là kết quả đã lưu từ đợt trước; không chạy lại toàn bộ bộ kiểm thử trong lượt chụp này.')])
    d.add_heading('Việc ưu tiên ở giai đoạn tiếp theo',2)
    for value in ['Tiếp tục kiểm thử dữ liệu biên, phân quyền, thay đổi tồn kho/giá và thao tác đặt lặp.',
      'Kiểm tra nhập tiếng Việt, bố cục ngang, cỡ chữ lớn, lưu hồ sơ và khôi phục sau khi đóng ứng dụng.',
      'Chạy trên Android API 28 và thiết bị dùng để nộp; kết quả hiện tại trên API 36 chưa thay thế kiểm tra phiên bản tối thiểu.',
      'Rà soát bộ kiểm thử, APK bàn giao và hoàn thiện tài liệu theo góp ý của giảng viên.',
      'Thống nhất phân công thực tế, lịch hoàn thành và thông tin lớp/giảng viên trước bản nộp chính thức.']:
        d.add_paragraph(value,style='List Bullet')
    d.add_paragraph('Giới hạn hiện tại: ứng dụng dùng SQLite trên một thiết bị, chưa có backend đồng bộ nhiều thiết bị, thanh toán trực tuyến hoặc tích hợp giao vận. Trạng thái Đang kiểm thử của dự án khác với trạng thái Đang giao của đơn hàng.')
    d.core_properties.title='Báo cáo tiến độ lần 1 – Bloomy Beauty – Nhóm 4'
    d.core_properties.author='Nhóm 4';d.core_properties.subject='Lập trình di động – 14/10/2026'
    target=OUT/'Bao-cao-tien-do-lan-1-Nhom-4.docx';d.save(target)
    manifest=[]
    for n,(key,title,desc) in enumerate(SCREENS,1):
        path=OUT/(key+'.png')
        with Image.open(path) as im:im.verify()
        manifest.append({'step':n,'file':path.name,'title':title,'description':desc,'captured_at_local':datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec='seconds'),'sha256':hashlib.sha256(path.read_bytes()).hexdigest()})
    (OUT/'danh-muc-anh.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
    print(target)
    print(f'{len(SCREENS)} images; {target.stat().st_size:,} bytes')

if __name__=='__main__':make()
